"""
core/doc_export.py
==================
One place that lets every "export as PDF" feature in the system also produce
a Word (.docx) document — the person clicking Export chooses PDF or Word.

How it works
────────────
1.  The browser adds ``?export_format=docx`` (or ``pdf``) to an export URL.
    ``ExportFormatMiddleware`` reads it once per request and stores it in a
    context variable, so any code running for that request can ask
    ``wants_docx()``.  Background threads / Celery tasks never inherit it,
    so anything that *stores* a document (published PDFs, caches, allocation
    runs) keeps producing real PDFs unless it is explicitly told otherwise.

2.  ReportLab builders
    ``SimpleDocTemplate`` and ``BaseDocTemplate`` exported from this module
    are drop-in replacements for the ReportLab classes.  In PDF mode they are
    exactly ReportLab.  In Word mode ``build()`` walks the very same list of
    flowables (Paragraph, Table, Spacer, Image, PageBreak, ...) and writes a
    .docx into the same buffer/filename instead — so the Word file always
    contains exactly what the PDF would have contained (same tables, same
    colours, same letterhead), with no second copy of any report logic.

3.  WeasyPrint (HTML → PDF) builders
    ``HTML`` exported from this module is a drop-in for ``weasyprint.HTML``.
    In Word mode ``write_pdf()`` converts the rendered HTML to .docx
    (headings, paragraphs, tables with colours/spans, lists, images).

4.  Already-stored PDFs (published timetables, cached per-program PDFs, ...)
    When Word is requested and the response is a finished PDF, the middleware
    converts it with ``pdf2docx`` (falling back to a pdfplumber text/table
    extraction if ``pdf2docx`` is not installed).  Conversions are cached on
    disk by content hash so the same file is only converted once.

5.  The middleware finally fixes the response headers (content-type,
    ``.pdf`` → ``.docx`` file name, attachment disposition, length).
"""
from __future__ import annotations

import base64
import contextvars
import hashlib
import io
import logging
import os
import re
import tempfile
from contextlib import contextmanager
from urllib.parse import unquote, urlparse

from django.conf import settings
from django.http import HttpResponse

logger = logging.getLogger(__name__)

FORMAT_PDF = "pdf"
FORMAT_DOCX = "docx"
FORMAT_PARAM = "export_format"

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

_current_format = contextvars.ContextVar("export_format", default=FORMAT_PDF)


# ═══════════════════════════════════════════════════════════════════════════
# Format selection
# ═══════════════════════════════════════════════════════════════════════════
def normalize_format(value) -> str:
    """'docx' / 'word' / 'doc' → docx; anything else → pdf."""
    v = str(value or "").strip().lower()
    return FORMAT_DOCX if v in ("docx", "word", "doc") else FORMAT_PDF


def current_format() -> str:
    return _current_format.get()


def wants_docx() -> bool:
    return _current_format.get() == FORMAT_DOCX


def requested_format(request, default=FORMAT_PDF) -> str:
    """Format asked for on a request (GET first, then POST)."""
    raw = None
    if request is not None:
        raw = request.GET.get(FORMAT_PARAM)
        if raw is None and getattr(request, "method", "") == "POST":
            raw = request.POST.get(FORMAT_PARAM)
    return normalize_format(raw) if raw is not None else default


@contextmanager
def force_format(fmt):
    token = _current_format.set(normalize_format(fmt))
    try:
        yield
    finally:
        _current_format.reset(token)


def force_pdf():
    """Use around code that STORES a document — it must always be a real PDF."""
    return force_format(FORMAT_PDF)


def extension_for(fmt=None) -> str:
    return ".docx" if normalize_format(fmt or current_format()) == FORMAT_DOCX else ".pdf"


def swap_extension(filename: str, fmt=None) -> str:
    """'report.pdf' → 'report.docx' when Word is active; unchanged otherwise."""
    if normalize_format(fmt or current_format()) != FORMAT_DOCX:
        return filename
    return re.sub(r"\.pdf$", ".docx", filename, flags=re.I) if re.search(r"\.pdf$", filename, re.I) else filename


# ═══════════════════════════════════════════════════════════════════════════
# Middleware
# ═══════════════════════════════════════════════════════════════════════════
class ExportFormatMiddleware:
    """
    Reads ``?export_format=`` and, for Word requests, post-processes any
    PDF response into a Word document (see module docstring, points 4 & 5).
    """

    def __init__(self, get_response):
        self.get_response = get_response

    def __call__(self, request):
        # Only plain GET/HEAD requests can switch the format from the URL.
        # POST endpoints (publish, generate, bulk actions) therefore always
        # keep producing real PDFs; the few POST exports that offer a Word
        # choice (feedback export, e-mail attachments) read the form field
        # themselves and call `force_format(...)` explicitly.
        fmt = FORMAT_PDF
        if request.method in ("GET", "HEAD"):
            fmt = normalize_format(request.GET.get(FORMAT_PARAM))
        token = _current_format.set(fmt)
        try:
            response = self.get_response(request)
            if fmt == FORMAT_DOCX:
                response = to_word_response(response)
            return response
        finally:
            _current_format.reset(token)


def to_word_response(response):
    """Turn a PDF response (native-.docx-in-a-PDF-envelope or a finished PDF)
    into a proper Word download.  Anything else is returned untouched."""
    try:
        if response.status_code != 200:
            return response
        ctype = (response.get("Content-Type") or "").lower()
        if not ctype.startswith("application/pdf"):
            return response
        if getattr(response, "streaming", False):
            body = b"".join(response.streaming_content)
        else:
            body = response.content
        if body[:2] == b"PK":                     # already a .docx (native build)
            docx_bytes = body
        elif body[:4] == b"%PDF":                 # a finished PDF → convert it
            docx_bytes = pdf_to_docx_bytes(body)
        else:
            return response
        return _docx_response_like(response, docx_bytes)
    except Exception:
        logger.exception("Word export failed — returning the PDF instead")
        return response



def _docx_response_like(old, docx_bytes: bytes):
    """Build the Word response, keeping the original file name / cookies."""
    disp = old.get("Content-Disposition", "") or ""
    m = re.search(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)"?', disp)
    filename = unquote(m.group(1)) if m else "document.pdf"
    filename = re.sub(r"\.pdf$", "", filename, flags=re.I) + ".docx"
    new = HttpResponse(docx_bytes, content_type=DOCX_CONTENT_TYPE)
    new["Content-Disposition"] = f'attachment; filename="{filename}"'
    new["Content-Length"] = str(len(docx_bytes))
    new["Cache-Control"] = "no-store"
    for name, morsel in old.cookies.items():
        new.cookies[name] = morsel
    return new


def docx_response(docx_bytes: bytes, filename: str) -> HttpResponse:
    """Direct helper for views that already hold .docx bytes."""
    filename = re.sub(r"\.(pdf|docx)$", "", filename, flags=re.I) + ".docx"
    resp = HttpResponse(docx_bytes, content_type=DOCX_CONTENT_TYPE)
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    resp["Content-Length"] = str(len(docx_bytes))
    return resp


# ═══════════════════════════════════════════════════════════════════════════
# Small helpers shared by both converters
# ═══════════════════════════════════════════════════════════════════════════
def _font_family(name: str) -> str:
    n = (name or "").lower()
    if "times" in n or ("serif" in n and "sans" not in n):
        return "Times New Roman"
    if "courier" in n or "mono" in n:
        return "Courier New"
    return "Arial"


def _hex_from_color(c):
    """ReportLab Color → 'RRGGBB' (None if transparent / unknown)."""
    if c is None:
        return None
    try:
        if getattr(c, "alpha", 1) == 0:
            return None
        r, g, b = c.red, c.green, c.blue
        return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))
    except Exception:
        return None


def _set_cell_shading(cell, hex_fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _set_par_shading(par, hex_fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def _set_cell_border(cell, edge, size_pt, hex_color="000000"):
    """edge in top/left/bottom/right; size in points."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = borders.find(qn(f"w:{edge}"))
    if el is None:
        el = OxmlElement(f"w:{edge}")
        borders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(max(2, int(round(float(size_pt) * 8)))))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), hex_color or "000000")


def _set_cell_margins(cell, top=None, left=None, bottom=None, right=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    mar = tcPr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcPr.append(mar)
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        if val is None:
            continue
        el = mar.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            mar.append(el)
        el.set(qn("w:w"), str(int(float(val) * 20)))
        el.set(qn("w:type"), "dxa")


def _set_valign(cell, valign):
    from docx.enum.table import WD_ALIGN_VERTICAL
    v = (valign or "").upper()
    cell.vertical_alignment = {
        "TOP": WD_ALIGN_VERTICAL.TOP,
        "MIDDLE": WD_ALIGN_VERTICAL.CENTER,
        "BOTTOM": WD_ALIGN_VERTICAL.BOTTOM,
    }.get(v, WD_ALIGN_VERTICAL.TOP)


def _repeat_header(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_split_row(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _fixed_layout(table, col_widths_pt=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    table.autofit = False
    if col_widths_pt:
        for idx, w in enumerate(col_widths_pt):
            if idx < len(table.columns):
                table.columns[idx].width = Pt(w)
        for row in table.rows:
            for idx, w in enumerate(col_widths_pt):
                if idx < len(row.cells):
                    row.cells[idx].width = Pt(w)


def _add_field(run, instr, placeholder="1"):
    """Insert a Word field (PAGE, NUMPAGES...) into one run, styled like the run."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = f" {instr} "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = placeholder
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, i, sep, t, e):
        run._r.append(el)


def _style_run(run, *, font=None, size=None, bold=None, italic=None,
               underline=None, strike=None, color=None, sup=False, sub=False):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    if font:
        run.font.name = font
        rpr = run._r.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            rf.set(qn("w:eastAsia"), font)
            rf.set(qn("w:cs"), font)
    if size:
        run.font.size = Pt(max(1.0, min(float(size), 400.0)))
    if bold is not None:
        run.font.bold = bool(bold)
    if italic is not None:
        run.font.italic = bool(italic)
    if underline:
        run.font.underline = True
    if strike:
        run.font.strike = True
    if color:
        try:
            run.font.color.rgb = RGBColor.from_string(color.upper())
        except Exception:
            pass
    if sup:
        run.font.superscript = True
    if sub:
        run.font.subscript = True


def _apply_page_setup(document, page_w_pt, page_h_pt, left_pt, right_pt, top_pt, bottom_pt):
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt
    sec = document.sections[0]
    landscape = page_w_pt > page_h_pt
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width = Pt(page_w_pt)
    sec.page_height = Pt(page_h_pt)
    sec.left_margin = Pt(left_pt)
    sec.right_margin = Pt(right_pt)
    sec.top_margin = Pt(top_pt)
    sec.bottom_margin = Pt(bottom_pt)
    sec.header_distance = Pt(min(top_pt / 2, 30))
    sec.footer_distance = Pt(min(bottom_pt / 2, 24))
    return sec


def _write_footer(document, lines):
    """Footer: any captured text lines (Compiled by, watermark code) + Page N."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    sec = document.sections[0]
    footer = sec.footer
    footer.is_linked_to_previous = False
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    text = "   |   ".join(x for x in lines if x)
    if text:
        r = par.add_run(text + "   |   ")
        _style_run(r, font="Arial", size=7.5, color="777777")
    r = par.add_run("Page ")
    _style_run(r, font="Arial", size=7.5, color="777777")
    r2 = par.add_run()
    _style_run(r2, font="Arial", size=7.5, color="777777")
    _add_field(r2, "PAGE")


# ═══════════════════════════════════════════════════════════════════════════
# 1) ReportLab flowables  →  .docx
# ═══════════════════════════════════════════════════════════════════════════
class _RecordingCanvas:
    """Stand-in canvas that just records every string a page callback would
    draw (footer text, 'Compiled by', the anti-forgery watermark code, ...)."""

    def __init__(self):
        self.strings = []

    def _rec(self, *args, **kw):
        for a in args:
            if isinstance(a, str):
                self.strings.append(a)

    drawString = drawRightString = drawCentredString = _rec

    def __getattr__(self, name):        # every other canvas call is a no-op
        return lambda *a, **k: None


def _capture_footer_lines(doc, on_first):
    if not on_first:
        return []
    try:
        canvas = _RecordingCanvas()
        try:
            doc.page = 1
        except Exception:
            pass
        on_first(canvas, doc)
        seen, out = set(), []
        for s in canvas.strings:
            s = (s or "").strip()
            if not s or s.lower().startswith("page ") or s in seen:
                continue
            seen.add(s)
            out.append(s)
        return out
    except Exception:
        return []


class FlowableDocxWriter:
    """Renders a ReportLab story into a python-docx Document."""

    def __init__(self, pagesize, left, right, top, bottom):
        from docx import Document
        self.document = Document()
        self.page_w, self.page_h = float(pagesize[0]), float(pagesize[1])
        self.left, self.right = float(left), float(right)
        self.avail = max(72.0, self.page_w - self.left - self.right)
        _apply_page_setup(self.document, self.page_w, self.page_h, left, right, top, bottom)
        st = self.document.styles["Normal"]
        st.paragraph_format.space_after = 0
        st.paragraph_format.space_before = 0
        self._pending_break = False

    # ── public ─────────────────────────────────────────────────────────
    def write(self, story):
        for fl in story:
            self._flowable(fl, self.document, self.avail)

    def save(self, target):
        self.document.save(target)

    # ── dispatch ───────────────────────────────────────────────────────
    def _flowable(self, fl, container, avail):
        from reportlab.platypus import (
            Paragraph, Spacer, Table, PageBreak, Image, HRFlowable,
        )
        try:
            from reportlab.platypus.flowables import (
                KeepTogether, KeepInFrame, CondPageBreak, Indenter,
            )
        except Exception:                                   # pragma: no cover
            KeepTogether = KeepInFrame = CondPageBreak = Indenter = ()

        if fl is None:
            return
        if isinstance(fl, (list, tuple)):
            for x in fl:
                self._flowable(x, container, avail)
        elif isinstance(fl, Paragraph):
            self._paragraph(fl, container)
        elif isinstance(fl, Table):
            self._table(fl, container, avail)
        elif isinstance(fl, Spacer):
            self._spacer(fl, container)
        elif isinstance(fl, PageBreak):
            self._pending_break = True
        elif isinstance(fl, Image):
            self._image(fl, container)
        elif isinstance(fl, HRFlowable):
            self._hr(fl, container)
        elif KeepTogether and isinstance(fl, (KeepTogether, KeepInFrame)):
            for x in getattr(fl, "_content", []) or []:
                self._flowable(x, container, avail)
        elif CondPageBreak and isinstance(fl, (CondPageBreak, Indenter)):
            return
        elif hasattr(fl, "getPlainText"):                   # Preformatted etc.
            self._plain(fl.getPlainText(), container)
        elif hasattr(fl, "_content"):
            for x in fl._content or []:
                self._flowable(x, container, avail)
        # anything else (pure drawing flowables, anchors, ...) is skipped

    # ── building blocks ────────────────────────────────────────────────
    def _new_par(self, container):
        """A fresh paragraph — or the empty first one of a fresh table cell."""
        from docx.table import _Cell
        if isinstance(container, _Cell):
            first = container.paragraphs[0]
            if not getattr(container, "_first_used", False) and not first.text and len(container.paragraphs) == 1:
                container._first_used = True
                par = first
            else:
                container._first_used = True
                par = container.add_paragraph()
        else:
            par = container.add_paragraph()
        if self._pending_break and not isinstance(container, _Cell):
            par.paragraph_format.page_break_before = True
            self._pending_break = False
        return par

    def _spacer(self, sp, container):
        from docx.shared import Pt
        h = float(getattr(sp, "height", 0) or 0)
        if h < 3:
            return
        par = self._new_par(container)
        pf = par.paragraph_format
        pf.space_before = 0
        pf.space_after = 0
        try:
            from docx.enum.text import WD_LINE_SPACING
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(min(h, 60))
        except Exception:
            pass
        _style_run(par.add_run(""), size=1)

    def _plain(self, text, container):
        par = self._new_par(container)
        _style_run(par.add_run(str(text or "")), font="Arial", size=9)

    def _hr(self, hr, container):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        par = self._new_par(container)
        par.paragraph_format.space_after = 4
        _style_run(par.add_run(""), size=2)
        pPr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(max(2, int(float(getattr(hr, "lineWidth", 1) or 1) * 8))))
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), _hex_from_color(getattr(hr, "color", None)) or "000000")
        bdr.append(bottom)
        pPr.append(bdr)

    def _image(self, img, container, par=None):
        from docx.shared import Pt
        src = getattr(img, "filename", None) or getattr(img, "_filename", None)
        stream = None
        if isinstance(src, str) and os.path.exists(src):
            stream = src
        elif hasattr(src, "read"):
            try:
                src.seek(0)
                stream = io.BytesIO(src.read())
            except Exception:
                stream = None
        if stream is None:
            return
        par = par or self._new_par(container)
        h_align = (getattr(img, "hAlign", "CENTER") or "CENTER").upper()
        self._align(par, {"CENTER": 1, "RIGHT": 2}.get(h_align, 0))
        try:
            width = float(getattr(img, "drawWidth", 0) or 0)
            par.add_run().add_picture(stream, width=Pt(width) if width else None)
        except Exception:
            logger.debug("docx: could not embed image", exc_info=True)

    @staticmethod
    def _align(par, code):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        par.alignment = {
            0: WD_ALIGN_PARAGRAPH.LEFT, 1: WD_ALIGN_PARAGRAPH.CENTER,
            2: WD_ALIGN_PARAGRAPH.RIGHT, 4: WD_ALIGN_PARAGRAPH.JUSTIFY,
        }.get(code, WD_ALIGN_PARAGRAPH.LEFT)

    def _paragraph(self, p, container):
        from docx.shared import Pt
        par = self._new_par(container)
        style = p.style
        self._align(par, getattr(style, "alignment", 0))
        pf = par.paragraph_format
        pf.space_before = Pt(float(getattr(style, "spaceBefore", 0) or 0))
        pf.space_after = Pt(float(getattr(style, "spaceAfter", 0) or 0))
        if getattr(style, "leftIndent", 0):
            pf.left_indent = Pt(float(style.leftIndent))
        if getattr(style, "firstLineIndent", 0):
            pf.first_line_indent = Pt(float(style.firstLineIndent))
        bg = _hex_from_color(getattr(style, "backColor", None))
        if bg:
            _set_par_shading(par, bg)

        style_font = getattr(style, "fontName", "Helvetica")
        frags = getattr(p, "frags", None)
        if not frags:
            frags = []
        wrote = False
        for fr in frags:
            if getattr(fr, "lineBreak", False):
                par.add_run().add_break()
                continue
            text = getattr(fr, "text", "") or ""
            if not text:
                continue
            fname = getattr(fr, "fontName", style_font) or style_font
            rise = getattr(fr, "rise", 0) or 0
            run = par.add_run(text)
            _style_run(
                run,
                font=_font_family(fname),
                size=getattr(fr, "fontSize", None) or getattr(style, "fontSize", 10),
                bold=("bold" in fname.lower()) or bool(getattr(fr, "bold", 0)),
                italic=("italic" in fname.lower() or "oblique" in fname.lower()) or bool(getattr(fr, "italic", 0)),
                underline=any(len(k) > 1 and k[1] == "underline" for k in (getattr(fr, "us_lines", None) or [])),
                strike=any(len(k) > 1 and k[1] == "strike" for k in (getattr(fr, "us_lines", None) or [])),
                color=_hex_from_color(getattr(fr, "textColor", None)),
                sup=rise > 0,
                sub=rise < 0,
            )
            wrote = True
        if not wrote:
            try:
                txt = p.getPlainText()
            except Exception:
                txt = ""
            if txt:
                _style_run(par.add_run(txt), font=_font_family(style_font),
                           size=getattr(style, "fontSize", 10),
                           bold="bold" in style_font.lower(),
                           color=_hex_from_color(getattr(style, "textColor", None)))

    # ── tables ─────────────────────────────────────────────────────────
    @staticmethod
    def _resolve_widths(arg_w, ncols, avail):
        widths = []
        for w in list(arg_w or [])[:ncols]:
            if isinstance(w, str) and w.strip().endswith("%"):
                try:
                    widths.append(avail * float(w.strip()[:-1]) / 100.0)
                except ValueError:
                    widths.append(None)
            elif isinstance(w, (int, float)):
                widths.append(float(w))
            else:
                widths.append(None)
        widths += [None] * (ncols - len(widths))
        known = sum(w for w in widths if w)
        missing = [i for i, w in enumerate(widths) if not w]
        if missing:
            share = max(24.0, (avail - known) / len(missing)) if avail > known else 48.0
            for i in missing:
                widths[i] = share
        total = sum(widths)
        if total > avail and total > 0:
            widths = [w * avail / total for w in widths]
        return widths

    @staticmethod
    def _rng(cmd, nrows, ncols):
        (sc, sr), (ec, er) = cmd[1], cmd[2]
        sc = sc + ncols if sc < 0 else sc
        ec = ec + ncols if ec < 0 else ec
        sr = sr + nrows if sr < 0 else sr
        er = er + nrows if er < 0 else er
        return max(0, min(sc, ec)), max(0, min(sr, er)), min(ncols - 1, max(sc, ec)), min(nrows - 1, max(sr, er))

    def _table(self, tbl, container, avail):
        from docx.table import _Cell
        data = getattr(tbl, "_cellvalues", None) or []
        nrows = len(data)
        if not nrows:
            return
        ncols = max(len(r) for r in data)
        if ncols == 0:
            return

        widths = self._resolve_widths(getattr(tbl, "_argW", None) or getattr(tbl, "_colWidths", None), ncols, avail)

        if isinstance(container, _Cell):
            dt = container.add_table(nrows, ncols)
        else:
            if self._pending_break:                         # honour a pending PageBreak
                par = container.add_paragraph()
                par.paragraph_format.page_break_before = True
                _style_run(par.add_run(""), size=1)
                self._pending_break = False
            dt = container.add_table(rows=nrows, cols=ncols)
        _fixed_layout(dt, widths)
        if not isinstance(container, _Cell) and str(getattr(tbl, "hAlign", "CENTER")).upper() == "CENTER":
            from docx.enum.table import WD_TABLE_ALIGNMENT
            dt.alignment = WD_TABLE_ALIGNMENT.CENTER

        span_map = set()                                     # cells hidden by SPAN
        spans = []
        for cmd in getattr(tbl, "_spanCmds", []) or []:
            c0, r0, c1, r1 = self._rng(cmd, nrows, ncols)
            spans.append((c0, r0, c1, r1))
            for rr in range(r0, r1 + 1):
                for cc in range(c0, c1 + 1):
                    if (cc, rr) != (c0, r0):
                        span_map.add((cc, rr))

        # ── backgrounds (later commands win, like ReportLab) ──
        fills = {}
        for cmd in getattr(tbl, "_bkgrndcmds", []) or []:
            op = cmd[0]
            c0, r0, c1, r1 = self._rng(cmd, nrows, ncols)
            if op == "BACKGROUND":
                hexv = _hex_from_color(cmd[3])
                for rr in range(r0, r1 + 1):
                    for cc in range(c0, c1 + 1):
                        fills[(cc, rr)] = hexv
            elif op == "ROWBACKGROUNDS":
                cols = [_hex_from_color(c) for c in (cmd[3] or [])]
                if cols:
                    for i, rr in enumerate(range(r0, r1 + 1)):
                        for cc in range(c0, c1 + 1):
                            fills[(cc, rr)] = cols[i % len(cols)]
            elif op == "COLBACKGROUNDS":
                cols = [_hex_from_color(c) for c in (cmd[3] or [])]
                if cols:
                    for i, cc in enumerate(range(c0, c1 + 1)):
                        for rr in range(r0, r1 + 1):
                            fills[(cc, rr)] = cols[i % len(cols)]

        # ── borders ──
        borders = {}                                         # (c,r) -> {edge: (pt, hex)}

        def _edge(c, r, edge, wt, col):
            borders.setdefault((c, r), {})[edge] = (wt, col)

        for cmd in getattr(tbl, "_linecmds", []) or []:
            op = cmd[0]
            c0, r0, c1, r1 = self._rng(cmd, nrows, ncols)
            wt, col = cmd[3], _hex_from_color(cmd[4]) or "000000"
            if wt is None or float(wt) <= 0:
                continue
            for rr in range(r0, r1 + 1):
                for cc in range(c0, c1 + 1):
                    if op == "GRID":
                        for e in ("top", "bottom", "left", "right"):
                            _edge(cc, rr, e, wt, col)
                    elif op in ("BOX", "OUTLINE"):
                        if rr == r0: _edge(cc, rr, "top", wt, col)
                        if rr == r1: _edge(cc, rr, "bottom", wt, col)
                        if cc == c0: _edge(cc, rr, "left", wt, col)
                        if cc == c1: _edge(cc, rr, "right", wt, col)
                    elif op == "INNERGRID":
                        if rr != r0: _edge(cc, rr, "top", wt, col)
                        if rr != r1: _edge(cc, rr, "bottom", wt, col)
                        if cc != c0: _edge(cc, rr, "left", wt, col)
                        if cc != c1: _edge(cc, rr, "right", wt, col)
                    elif op == "LINEBELOW":
                        _edge(cc, rr, "bottom", wt, col)
                    elif op == "LINEABOVE":
                        _edge(cc, rr, "top", wt, col)
                    elif op == "LINEBEFORE":
                        _edge(cc, rr, "left", wt, col)
                    elif op == "LINEAFTER":
                        _edge(cc, rr, "right", wt, col)

        # ── repeat header rows ──
        rep = getattr(tbl, "repeatRows", 0)
        rep_rows = list(rep) if isinstance(rep, (list, tuple)) else list(range(int(rep or 0)))

        for r in range(nrows):
            row_vals = data[r]
            if r in rep_rows:
                _repeat_header(dt.rows[r])
            _no_split_row(dt.rows[r])
            for c in range(ncols):
                cell = dt.cell(r, c)
                cs = None
                try:
                    cs = tbl._cellStyles[r][c]
                except Exception:
                    pass
                if cs is not None:
                    _set_valign(cell, getattr(cs, "valign", "TOP"))
                    _set_cell_margins(cell, getattr(cs, "topPadding", None), getattr(cs, "leftPadding", None),
                                      getattr(cs, "bottomPadding", None), getattr(cs, "rightPadding", None))
                if fills.get((c, r)):
                    _set_cell_shading(cell, fills[(c, r)])
                for edge, (wt, col) in borders.get((c, r), {}).items():
                    _set_cell_border(cell, edge, wt, col)
                if (c, r) in span_map or c >= len(row_vals):
                    continue
                self._cell_content(row_vals[c], cell, widths[c] if c < len(widths) else 60.0, cs)

        for c0, r0, c1, r1 in spans:
            try:
                dt.cell(r0, c0).merge(dt.cell(r1, c1))
            except Exception:
                logger.debug("docx: SPAN merge failed", exc_info=True)

        if not isinstance(container, _Cell):
            # a tiny paragraph after a table keeps consecutive tables from fusing
            gap = container.add_paragraph()
            gap.paragraph_format.space_after = 0
            _style_run(gap.add_run(""), size=2)

    def _cell_content(self, val, cell, width, cs):
        from reportlab.platypus import Paragraph, Table
        from docx.shared import Pt
        if val is None or val == "":
            return
        if isinstance(val, (str, int, float)):
            par = self._new_par(cell)
            align = (getattr(cs, "alignment", "LEFT") or "LEFT")
            if isinstance(align, str):
                self._align(par, {"LEFT": 0, "CENTER": 1, "CENTRE": 1, "RIGHT": 2}.get(align.upper(), 0))
            fname = getattr(cs, "fontname", "Helvetica") or "Helvetica"
            _style_run(
                par.add_run(str(val)),
                font=_font_family(fname), size=getattr(cs, "fontsize", 10) or 10,
                bold="bold" in fname.lower(),
                italic="italic" in fname.lower() or "oblique" in fname.lower(),
                color=_hex_from_color(getattr(cs, "color", None)),
            )
            return
        inner = max(24.0, width - 8.0)
        self._flowable(val, cell, inner)


def flowables_to_docx_bytes(story, *, pagesize, left, right, top, bottom, footer_lines=None) -> bytes:
    w = FlowableDocxWriter(pagesize, left, right, top, bottom)
    w.write(story)
    _write_footer(w.document, footer_lines or [])
    buf = io.BytesIO()
    w.save(buf)
    return buf.getvalue()


# ─── drop-in ReportLab document templates ───────────────────────────────
from reportlab.platypus import SimpleDocTemplate as _RLSimpleDocTemplate  # noqa: E402
from reportlab.platypus import BaseDocTemplate as _RLBaseDocTemplate      # noqa: E402


def _build_docx_into(doc, flowables, on_first):
    """Write `flowables` as .docx into whatever `doc.filename` points at."""
    story = list(flowables)
    footer_lines = list(getattr(doc, "docx_footer_lines", None) or [])
    footer_lines += _capture_footer_lines(doc, on_first)
    data = flowables_to_docx_bytes(
        story,
        pagesize=doc.pagesize,
        left=getattr(doc, "leftMargin", 36), right=getattr(doc, "rightMargin", 36),
        top=getattr(doc, "topMargin", 36), bottom=getattr(doc, "bottomMargin", 36),
        footer_lines=footer_lines,
    )
    target = doc.filename
    if hasattr(target, "write"):
        target.write(data)
    else:
        with open(target, "wb") as fh:
            fh.write(data)


class SimpleDocTemplate(_RLSimpleDocTemplate):
    """ReportLab's SimpleDocTemplate — that writes .docx when Word is requested."""

    def build(self, flowables, onFirstPage=None, onLaterPages=None, *args, **kwargs):
        if wants_docx():
            _build_docx_into(self, flowables, onFirstPage)
            return
        if onFirstPage is None and onLaterPages is None and not args:
            return super().build(flowables, **kwargs)
        kw = dict(kwargs)
        if onFirstPage is not None:
            kw["onFirstPage"] = onFirstPage
        if onLaterPages is not None:
            kw["onLaterPages"] = onLaterPages
        return super().build(flowables, *args, **kw)


class BaseDocTemplate(_RLBaseDocTemplate):
    """ReportLab's BaseDocTemplate — that writes .docx when Word is requested."""

    def build(self, flowables, *args, **kwargs):
        if wants_docx():
            try:
                first = self.pageTemplates[0].onPage if self.pageTemplates else None
            except Exception:
                first = None
            _build_docx_into(self, flowables, first)
            return
        return super().build(flowables, *args, **kwargs)


# ═══════════════════════════════════════════════════════════════════════════
# 2) HTML (WeasyPrint-style)  →  .docx
# ═══════════════════════════════════════════════════════════════════════════
_NAMED_COLORS = {
    "black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000",
    "blue": "0000FF", "gray": "808080", "grey": "808080", "silver": "C0C0C0",
    "yellow": "FFFF00", "orange": "FFA500", "navy": "000080", "maroon": "800000",
    "purple": "800080", "teal": "008080", "lightgray": "D3D3D3", "lightgrey": "D3D3D3",
    "darkgray": "A9A9A9", "darkgrey": "A9A9A9",
}
_BLOCK_TAGS = {
    "address", "article", "aside", "blockquote", "body", "dd", "div", "dl", "dt",
    "fieldset", "figure", "footer", "form", "h1", "h2", "h3", "h4", "h5", "h6",
    "header", "hr", "li", "main", "nav", "ol", "p", "pre", "section", "table",
    "ul", "thead", "tbody", "tfoot", "tr", "caption",
}
_SKIP_TAGS = {"script", "style", "head", "title", "meta", "link", "noscript", "svg", "canvas", "button", "input", "select", "textarea"}
_HEADING_PT = {"h1": 20, "h2": 16, "h3": 14, "h4": 12, "h5": 11, "h6": 10}


def _css_color(value):
    """CSS colour value → 'RRGGBB' or None. Picks first colour inside gradients."""
    if not value:
        return None
    v = value.strip().lower()
    if v in ("transparent", "none", "inherit", "initial", "currentcolor"):
        return None
    m = re.search(r"#([0-9a-f]{6})\b", v)
    if m:
        return m.group(1).upper()
    m = re.search(r"#([0-9a-f]{3})\b", v)
    if m:
        return "".join(ch * 2 for ch in m.group(1)).upper()
    m = re.search(r"rgba?\(\s*(\d+)[,\s]+(\d+)[,\s]+(\d+)(?:[,\s/]+([\d.]+))?", v)
    if m:
        if m.group(4) is not None and float(m.group(4)) < 0.5:
            return None
        return "%02X%02X%02X" % tuple(min(255, int(m.group(i))) for i in (1, 2, 3))
    for name, hx in _NAMED_COLORS.items():
        if re.search(rf"\b{name}\b", v):
            return hx
    return None


def _css_len_pt(value, default=None):
    if not value:
        return default
    m = re.match(r"\s*(-?[\d.]+)\s*(px|pt|em|rem|mm|cm|in|%)?", str(value).lower())
    if not m:
        return default
    n = float(m.group(1))
    unit = m.group(2) or "px"
    return {"px": n * 0.75, "pt": n, "em": n * 10, "rem": n * 10,
            "mm": n * 2.835, "cm": n * 28.35, "in": n * 72, "%": default}.get(unit, default)


def _parse_declarations(text):
    out = {}
    for decl in (text or "").split(";"):
        if ":" in decl:
            k, v = decl.split(":", 1)
            out[k.strip().lower()] = v.strip()
    return out


def _collect_css_rules(soup):
    """Very small CSS reader: `sel, sel { decls }` with simple selectors."""
    rules = []
    for st in soup.find_all("style"):
        css = re.sub(r"/\*.*?\*/", "", st.get_text() or "", flags=re.S)
        css = re.sub(r"@page\s*\{(?:[^{}]|\{[^{}]*\})*\}", "", css, flags=re.S)
        css = re.sub(r"@media\s+print\s*\{", "", css)            # keep print rules
        for m in re.finditer(r"([^{}@]+)\{([^{}]*)\}", css):
            decls = _parse_declarations(m.group(2))
            for sel in m.group(1).split(","):
                sel = sel.strip()
                if sel and "::" not in sel and ":hover" not in sel:
                    rules.append((sel, decls))
    return rules


def _nth_matches(expr, idx):
    """:nth-child(expr) for odd / even / N / an+b."""
    e = expr.replace(" ", "").lower()
    if e == "odd":
        return idx % 2 == 1
    if e == "even":
        return idx % 2 == 0
    m = re.fullmatch(r"([+-]?\d*)n([+-]\d+)?", e)
    if m:
        a = m.group(1)
        a = -1 if a == "-" else (1 if a in ("", "+") else int(a))
        b = int(m.group(2) or 0)
        if a == 0:
            return idx == b
        return (idx - b) % a == 0 and (idx - b) // a >= 0
    return e.isdigit() and idx == int(e)


def _compound_matches(comp, tag):
    m = re.match(r"^([a-zA-Z][\w-]*|\*)?((?:[.#][\w-]+)*)((?::[\w-]+(?:\([^)]*\))?)*)$", comp)
    if not m:
        return False
    name, quals, pseudos = m.groups()
    if name and name != "*" and tag.name != name.lower():
        return False
    classes = tag.get("class") or []
    for q in re.findall(r"[.#][\w-]+", quals or ""):
        if q[0] == "." and q[1:] not in classes:
            return False
        if q[0] == "#" and tag.get("id") != q[1:]:
            return False
    if not (name or quals):
        return False
    for ps in re.findall(r":([\w-]+)(?:\(([^)]*)\))?", pseudos or ""):
        kind, arg = ps
        if kind == "first-child":
            if tag.find_previous_sibling(True) is not None:
                return False
        elif kind == "last-child":
            if tag.find_next_sibling(True) is not None:
                return False
        elif kind == "nth-child":
            idx = 1 + len(tag.find_previous_siblings(True))
            if not _nth_matches(arg, idx):
                return False
        else:
            return False            # :hover, :not(), ... — not applicable to a static document
    return True


def _sel_matches(sel, tag):
    """Descendant (' ') and child ('>') selectors with class/id/tag/nth-child."""
    tokens = re.findall(r"[>+~]|[^\s>+~]+", sel.strip())
    if not tokens or any(t in ("+", "~") for t in tokens):
        return False

    def match_from(idx, node):
        # tokens[idx] is a compound that must match `node`
        if not _compound_matches(tokens[idx], node):
            return False
        if idx == 0:
            return True
        comb = tokens[idx - 1] if tokens[idx - 1] == ">" else " "
        j = idx - 2 if tokens[idx - 1] == ">" else idx - 1
        if j < 0:
            return False
        if comb == ">":
            par = node.parent
            return par is not None and getattr(par, "name", None) not in (None, "[document]") and match_from(j, par)
        anc = node.parent
        while anc is not None and getattr(anc, "name", None) not in (None, "[document]"):
            if match_from(j, anc):
                return True
            anc = anc.parent
        return False

    return match_from(len(tokens) - 1, tag)


class _HtmlDocxWriter:
    def __init__(self, soup, base_url=""):
        from docx import Document
        self.soup = soup
        self.base_url = base_url or ""
        self.rules = _collect_css_rules(soup)
        self.document = Document()
        self._style_cache = {}
        self._pending_break = False
        st = self.document.styles["Normal"]
        st.paragraph_format.space_after = 0
        st.paragraph_format.space_before = 0

        page_w, page_h, margins = self._page_from_css()
        self.margins = margins
        _apply_page_setup(self.document, page_w, page_h, margins[3], margins[1], margins[0], margins[2])
        self.avail = page_w - margins[3] - margins[1]

    # ── page ──────────────────────────────────────────────────────────
    def _page_from_css(self):
        css = " ".join(st.get_text() for st in self.soup.find_all("style"))
        m = re.search(r"@page\s*\{((?:[^{}]|\{[^{}]*\})*)\}", css, flags=re.S)
        w, h = 595.28, 841.89
        mt = mr = mb = ml = 45.0
        if m:
            block = re.sub(r"@[\w-]+\s*\{[^{}]*\}", "", m.group(1))
            decl = _parse_declarations(block)
            size = decl.get("size", "").lower()
            if "a3" in size:
                w, h = 841.89, 1190.55
            if "letter" in size:
                w, h = 612, 792
            if "landscape" in size:
                w, h = max(w, h), min(w, h)
            if "margin" in decl:
                parts = [_css_len_pt(x, 45.0) for x in decl["margin"].split()]
                if len(parts) == 1: mt = mr = mb = ml = parts[0]
                elif len(parts) == 2: mt = mb = parts[0]; mr = ml = parts[1]
                elif len(parts) == 3: mt, mr, mb = parts; ml = mr
                elif len(parts) >= 4: mt, mr, mb, ml = parts[:4]
        return w, h, (mt, mr, mb, ml)

    # ── style resolution ──────────────────────────────────────────────
    def _style_of(self, tag, parent):
        key = id(tag)
        if key in self._style_cache:
            return self._style_cache[key]
        st = {}
        for k in ("color", "font-weight", "font-style", "font-size", "text-align",
                  "font-family", "text-decoration"):          # inherited props
            if k in parent:
                st[k] = parent[k]
        if tag.name in ("b", "strong", "th") or tag.name in _HEADING_PT:
            st["font-weight"] = "bold"
        if tag.name in ("i", "em"):
            st["font-style"] = "italic"
        if tag.name in ("u", "ins"):
            st["text-decoration"] = "underline"
        if tag.name in ("s", "del", "strike"):
            st["text-decoration"] = "line-through"
        if tag.name in _HEADING_PT:
            st["font-size"] = f"{_HEADING_PT[tag.name]}pt"
        if tag.name == "th":
            st.setdefault("text-align", "center")
        if tag.get("align"):
            st["text-align"] = tag.get("align")
        if tag.name == "font" and tag.get("color"):
            st["color"] = tag.get("color")
        if tag.get("bgcolor"):
            st["background-color"] = tag.get("bgcolor")
        own = {}
        for sel, decls in self.rules:
            if _sel_matches(sel, tag):
                own.update(decls)
        own.update(_parse_declarations(tag.get("style", "")))
        # shorthands
        if "background" in own and "background-color" not in own:
            own["background-color"] = own["background"]
        st.update(own)
        self._style_cache[key] = st
        return st

    @staticmethod
    def _is_hidden(st):
        return (st.get("display", "").strip().lower() == "none"
                or st.get("visibility", "").strip().lower() == "hidden")

    @staticmethod
    def _wants_break_before(tag, st):
        v = (st.get("page-break-before", "") + " " + st.get("break-before", "")).lower()
        cls = " ".join(tag.get("class") or []).lower()
        return "always" in v or "page" in v.split() or "page-break" in cls or "pagebreak" in cls

    # ── entry ─────────────────────────────────────────────────────────
    def render(self):
        body = self.soup.body or self.soup
        self._blocks(body, self.document, {}, self.avail, bg=None)
        return self.document

    # ── block walking ─────────────────────────────────────────────────
    def _blocks(self, node, container, inherited, avail, bg):
        pending = []                                         # inline nodes awaiting a paragraph

        def flush():
            if any((getattr(n, "strip", None) and n.strip()) or getattr(n, "name", None) for n in pending):
                self._inline_paragraph(pending[:], container, inherited, bg)
            pending.clear()

        from bs4 import NavigableString, Comment
        for child in node.children:
            if isinstance(child, Comment):
                continue
            if isinstance(child, NavigableString):
                if str(child).strip():
                    pending.append(child)
                elif pending:
                    pending.append(child)
                continue
            name = child.name
            if name in _SKIP_TAGS:
                continue
            st = self._style_of(child, inherited)
            if self._is_hidden(st):
                continue
            if name in _BLOCK_TAGS:
                flush()
                if self._wants_break_before(child, st):
                    self._pending_break = True
                self._block(child, container, st, avail, bg)
            elif name == "br" and not pending:
                continue
            elif name == "img" and not pending:
                flush()
                self._image_par(child, container, st)
            else:
                pending.append(child)
        flush()

    def _block(self, tag, container, st, avail, bg):
        name = tag.name
        new_bg = _css_color(st.get("background-color")) or bg
        if name == "table":
            self._table(tag, container, st, avail)
        elif name in ("ul", "ol"):
            self._list(tag, container, st, avail, new_bg)
        elif name == "hr":
            self._hr(container)
        elif name in ("thead", "tbody", "tfoot", "tr"):
            self._blocks(tag, container, st, avail, new_bg)
        elif name == "pre":
            par = self._new_par(container, st, new_bg)
            self._add_text_run(par, tag.get_text(), st, force_mono=True)
        else:
            has_block_child = any(getattr(c, "name", None) in _BLOCK_TAGS for c in tag.children)
            if has_block_child or name in ("div", "section", "body", "main", "article", "header", "footer", "form", "figure"):
                if has_block_child:
                    self._blocks(tag, container, st, avail, new_bg)
                else:
                    self._inline_paragraph(list(tag.children), container, st, new_bg, is_block=True, tag=tag)
            else:
                self._inline_paragraph(list(tag.children), container, st, new_bg, is_block=True, tag=tag)

    # ── paragraphs / runs ─────────────────────────────────────────────
    def _new_par(self, container, st, bg=None):
        from docx.table import _Cell
        from docx.shared import Pt
        if isinstance(container, _Cell):
            first = container.paragraphs[0]
            if not getattr(container, "_first_used", False) and not first.text and len(container.paragraphs) == 1:
                par = first
            else:
                par = container.add_paragraph()
            container._first_used = True
        else:
            par = container.add_paragraph()
            if self._pending_break:
                par.paragraph_format.page_break_before = True
                self._pending_break = False
        align = (st.get("text-align") or "").lower()
        from docx.enum.text import WD_ALIGN_PARAGRAPH as A
        par.alignment = {"center": A.CENTER, "right": A.RIGHT, "justify": A.JUSTIFY}.get(align, A.LEFT)
        mb = _css_len_pt(st.get("margin-bottom"), None)
        mt = _css_len_pt(st.get("margin-top"), None)
        if mb is not None:
            par.paragraph_format.space_after = Pt(min(mb, 36))
        if mt is not None:
            par.paragraph_format.space_before = Pt(min(mt, 36))
        if bg:
            _set_par_shading(par, bg)
        return par

    def _add_text_run(self, par, text, st, force_mono=False):
        if text is None or text == "":
            return
        weight = str(st.get("font-weight", "")).lower()
        bold = weight in ("bold", "bolder") or (weight.isdigit() and int(weight) >= 600)
        deco = st.get("text-decoration", "").lower()
        fam = "Courier New" if force_mono else _font_family(st.get("font-family", "Arial"))
        run = par.add_run(text)
        _style_run(
            run, font=fam,
            size=_css_len_pt(st.get("font-size"), 10) or 10,
            bold=bold, italic=("italic" in st.get("font-style", "").lower()),
            underline="underline" in deco, strike="line-through" in deco,
            color=_css_color(st.get("color")),
        )

    def _inline_paragraph(self, nodes, container, st, bg, is_block=False, tag=None):
        from bs4 import NavigableString
        if is_block and tag is not None and self._wants_break_before(tag, st):
            self._pending_break = True
        par = None
        buf_text = []

        def ensure():
            nonlocal par
            if par is None:
                par = self._new_par(container, st, bg)
            return par

        def emit(node, cur_st):
            if isinstance(node, NavigableString):
                txt = re.sub(r"\s+", " ", str(node))
                if txt:
                    p = ensure()
                    if not p.text and not p.runs:
                        txt = txt.lstrip()
                    self._add_text_run(p, txt, cur_st)
                return
            if getattr(node, "name", None) is None or node.name in _SKIP_TAGS:
                return
            child_st = self._style_of(node, cur_st)
            if self._is_hidden(child_st):
                return
            if node.name == "br":
                ensure().add_run().add_break()
            elif node.name == "img":
                self._image_into(ensure(), node, child_st)
            elif node.name in _BLOCK_TAGS:                   # block inside an "inline" run — flatten
                for c in node.children:
                    emit(c, child_st)
                ensure().add_run().add_break()
            else:
                for c in node.children:
                    emit(c, child_st)

        is_flex = "flex" in str(st.get("display", "")).lower()
        first_child = True
        for n in nodes:
            if is_flex and getattr(n, "name", None) and not first_child:
                emit(NavigableString("   \u2022   "), st)          # flex children sit side by side
            if getattr(n, "name", None):
                first_child = False
            emit(n, st)
        if par is not None:
            # trim trailing blank
            if par.runs and par.runs[-1].text:
                par.runs[-1].text = par.runs[-1].text.rstrip()

    # ── images ────────────────────────────────────────────────────────
    def _resolve_image(self, src):
        if not src:
            return None
        if src.startswith("data:"):
            try:
                head, b64 = src.split(",", 1)
                return io.BytesIO(base64.b64decode(b64)) if ";base64" in head else io.BytesIO(unquote(b64).encode())
            except Exception:
                return None
        path = urlparse(src).path if "://" in src else src
        if src.startswith("file://"):
            return path if os.path.exists(path) else None
        path = unquote(path)
        cands = []
        media_url = getattr(settings, "MEDIA_URL", "/media/") or "/media/"
        static_url = getattr(settings, "STATIC_URL", "/static/") or "/static/"
        if getattr(settings, "MEDIA_ROOT", None) and path.startswith(media_url):
            cands.append(os.path.join(str(settings.MEDIA_ROOT), path[len(media_url):]))
        if path.startswith(static_url):
            rel = path[len(static_url):]
            if getattr(settings, "STATIC_ROOT", None):
                cands.append(os.path.join(str(settings.STATIC_ROOT), rel))
            for d in getattr(settings, "STATICFILES_DIRS", []) or []:
                cands.append(os.path.join(str(d), rel))
        cands.append(path)
        for c in cands:
            if c and os.path.isfile(c):
                return c
        return None

    def _image_into(self, par, tag, st):
        from docx.shared import Pt
        img = self._resolve_image(tag.get("src", ""))
        if img is None:
            return
        w = _css_len_pt(tag.get("width") or st.get("width"), None)
        try:
            par.add_run().add_picture(img, width=Pt(min(w, self.avail)) if w else Pt(60))
        except Exception:
            logger.debug("docx: html image skipped", exc_info=True)

    def _image_par(self, tag, container, st):
        par = self._new_par(container, st)
        if not st.get("text-align"):
            from docx.enum.text import WD_ALIGN_PARAGRAPH as A
            par.alignment = A.CENTER
        self._image_into(par, tag, st)

    def _hr(self, container):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        par = self._new_par(container, {})
        _style_run(par.add_run(""), size=2)
        pPr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "1"); b.set(qn("w:color"), "808080")
        bdr.append(b)
        pPr.append(bdr)

    # ── lists ─────────────────────────────────────────────────────────
    def _list(self, tag, container, st, avail, bg):
        from docx.shared import Pt
        ordered = tag.name == "ol"
        n = 0
        for li in tag.find_all("li", recursive=False):
            n += 1
            lst = self._style_of(li, st)
            par = self._new_par(container, lst, bg)
            par.paragraph_format.left_indent = Pt(18)
            par.paragraph_format.first_line_indent = Pt(-12)
            self._add_text_run(par, (f"{n}. " if ordered else "\u2022 "), lst)
            from bs4 import NavigableString
            for c in li.children:
                if getattr(c, "name", None) in ("ul", "ol"):
                    self._list(c, container, self._style_of(c, lst), avail, bg)
                elif isinstance(c, NavigableString):
                    self._add_text_run(par, re.sub(r"\s+", " ", str(c)), lst)
                elif c.name == "br":
                    par.add_run().add_break()
                elif c.name not in _SKIP_TAGS:
                    self._add_text_run(par, re.sub(r"\s+", " ", c.get_text()), self._style_of(c, lst))

    # ── tables ────────────────────────────────────────────────────────
    def _table(self, tag, container, st, avail):
        from docx.table import _Cell
        rows = []
        for tr in tag.find_all("tr"):
            if tr.find_parent("table") is not tag:
                continue                                     # belongs to a nested table
            cells = [c for c in tr.find_all(["td", "th"], recursive=False)]
            if cells:
                rows.append((tr, cells))
        if not rows:
            return
        # occupancy grid with rowspan/colspan
        grid, placements, ncols = {}, [], 0
        for r, (tr, cells) in enumerate(rows):
            c = 0
            for cell in cells:
                while (r, c) in grid:
                    c += 1
                cs = max(1, int(re.sub(r"\D", "", cell.get("colspan", "1")) or 1))
                rs = max(1, int(re.sub(r"\D", "", cell.get("rowspan", "1")) or 1))
                for rr in range(r, min(len(rows), r + rs)):
                    for cc in range(c, c + cs):
                        grid[(rr, cc)] = True
                placements.append((r, c, rs, cs, cell, tr))
                c += cs
                ncols = max(ncols, c)
        nrows = len(rows)

        if isinstance(container, _Cell):
            dt = container.add_table(nrows, ncols)
        else:
            if self._pending_break:
                par = container.add_paragraph()
                par.paragraph_format.page_break_before = True
                _style_run(par.add_run(""), size=1)
                self._pending_break = False
            dt = container.add_table(rows=nrows, cols=ncols)

        # widths from first-row cell width hints, else equal
        hints = [None] * ncols
        for (r, c, rs, cs, cell, tr) in placements:
            if r == 0 and cs == 1:
                w = cell.get("width") or _parse_declarations(cell.get("style", "")).get("width") \
                    or self._style_of(cell, st).get("width")
                if w:
                    w = str(w).strip()
                    hints[c] = avail * float(w[:-1]) / 100 if w.endswith("%") and re.match(r"^[\d.]+%$", w) else _css_len_pt(w, None)
        widths = FlowableDocxWriter._resolve_widths(hints, ncols, avail)
        _fixed_layout(dt, widths)

        table_st = self._style_of(tag, st)
        table_border = bool(tag.get("border") and tag.get("border") != "0") or self._has_border(table_st)

        header_rows = set()
        for r, (tr, cells) in enumerate(rows):
            if tr.find_parent("thead") is not None:
                header_rows.add(r)
        for r in sorted(header_rows):
            _repeat_header(dt.rows[r])

        for (r, c, rs, cs, cell, tr) in placements:
            dc = dt.cell(r, c)
            cst = self._style_of(cell, self._style_of(tr, table_st))
            bgc = _css_color(cst.get("background-color")) or _css_color(self._style_of(tr, table_st).get("background-color"))
            if bgc:
                _set_cell_shading(dc, bgc)
            if table_border or self._has_border(cst):
                for e in ("top", "bottom", "left", "right"):
                    _set_cell_border(dc, e, 0.5, _css_color(cst.get("border-color")) or _border_hex(cst) or "999999")
            else:
                for e, key in (("bottom", "border-bottom"), ("top", "border-top"),
                               ("left", "border-left"), ("right", "border-right")):
                    if key in cst and _css_len_pt(cst[key].split()[0], 0) and "none" not in cst[key]:
                        _set_cell_border(dc, e, max(0.5, _css_len_pt(cst[key].split()[0], 0.5)), _css_color(cst[key]) or "999999")
            va = (cst.get("vertical-align") or cell.get("valign") or "top").lower()
            _set_valign(dc, {"middle": "MIDDLE", "bottom": "BOTTOM"}.get(va, "TOP"))
            pad = _css_len_pt(cst.get("padding", "").split()[0], None) if cst.get("padding") else None
            _set_cell_margins(dc, pad if pad is not None else 2, (pad if pad is not None else 4), pad if pad is not None else 2, (pad if pad is not None else 4))
            cell_w = sum(widths[c:c + cs]) if c < len(widths) else avail / max(1, ncols)
            self._blocks(cell, dc, cst, max(24.0, cell_w - 8), bg=None)

        for (r, c, rs, cs, cell, tr) in placements:
            if rs > 1 or cs > 1:
                try:
                    dt.cell(r, c).merge(dt.cell(min(nrows - 1, r + rs - 1), min(ncols - 1, c + cs - 1)))
                except Exception:
                    logger.debug("docx: html merge failed", exc_info=True)

        if not isinstance(container, _Cell):
            gap = container.add_paragraph()
            _style_run(gap.add_run(""), size=4)

    @staticmethod
    def _has_border(st):
        b = st.get("border", "")
        return bool(b) and "none" not in b.lower() and not b.strip().startswith("0")


def _border_hex(st):
    return _css_color(st.get("border", ""))


def html_to_docx_bytes(html_string: str, base_url: str = "") -> bytes:
    from bs4 import BeautifulSoup
    try:
        soup = BeautifulSoup(html_string, "lxml")
    except Exception:
        soup = BeautifulSoup(html_string, "html.parser")
    w = _HtmlDocxWriter(soup, base_url=base_url)
    doc = w.render()
    _write_footer(doc, [])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class HTML:
    """Drop-in for ``weasyprint.HTML`` — ``write_pdf()`` yields a .docx when Word is requested.

    Only the parts of the WeasyPrint API this project uses are supported
    (``string=``, ``base_url=``, ``write_pdf(target=None, stylesheets=...)``).
    """

    def __init__(self, string=None, base_url=None, filename=None, url=None, **kwargs):
        self._string = string
        self._base_url = base_url
        self._filename = filename
        self._url = url
        self._kwargs = kwargs

    def _real(self):
        from weasyprint import HTML as _WeasyHTML
        return _WeasyHTML(string=self._string, base_url=self._base_url,
                          filename=self._filename, url=self._url, **self._kwargs)

    def write_pdf(self, target=None, *args, **kwargs):
        if wants_docx():
            html = self._string
            if html is None and self._filename:
                with open(self._filename, encoding="utf-8") as fh:
                    html = fh.read()
            data = html_to_docx_bytes(html or "", base_url=self._base_url or "")
            if target is None:
                return data
            if hasattr(target, "write"):
                target.write(data)
            else:
                with open(target, "wb") as fh:
                    fh.write(data)
            return None
        return self._real().write_pdf(target, *args, **kwargs)

    def __getattr__(self, item):                              # render(), etc. → real WeasyPrint
        return getattr(self._real(), item)


# ═══════════════════════════════════════════════════════════════════════════
# 3) Finished PDF bytes → .docx  (stored / cached / published PDFs)
# ═══════════════════════════════════════════════════════════════════════════
def _docx_cache_dir():
    base = getattr(settings, "MEDIA_ROOT", None) or tempfile.gettempdir()
    path = os.path.join(str(base), "word_export_cache")
    try:
        os.makedirs(path, exist_ok=True)
        return path
    except OSError:
        return tempfile.gettempdir()


def pdf_to_docx_bytes(pdf_bytes: bytes) -> bytes:
    """
    Convert a finished PDF into an editable Word document.
    Uses ``pdf2docx`` when installed (best layout fidelity) and falls back to a
    pdfplumber text + table extraction.  Results are cached by content hash.
    """
    digest = hashlib.sha256(pdf_bytes).hexdigest()
    cache_path = os.path.join(_docx_cache_dir(), f"{digest}.docx")
    if os.path.isfile(cache_path):
        try:
            with open(cache_path, "rb") as fh:
                return fh.read()
        except OSError:
            pass

    data = None
    try:
        from pdf2docx import Converter                        # type: ignore
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "out.docx")
            cv = Converter(stream=pdf_bytes)
            try:
                cv.convert(out)
            finally:
                cv.close()
            with open(out, "rb") as fh:
                data = fh.read()
    except ImportError:
        logger.info("pdf2docx not installed — using the pdfplumber fallback for PDF → Word")
    except Exception:
        logger.exception("pdf2docx conversion failed — using the pdfplumber fallback")

    if not data:
        data = _pdfplumber_to_docx(pdf_bytes)

    try:
        with open(cache_path, "wb") as fh:
            fh.write(data)
    except OSError:
        pass
    return data


def _pdfplumber_to_docx(pdf_bytes: bytes) -> bytes:
    import pdfplumber
    from docx import Document
    from docx.shared import Pt
    from docx.enum.section import WD_ORIENT

    doc = Document()
    doc.styles["Normal"].paragraph_format.space_after = Pt(2)
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for idx, page in enumerate(pdf.pages):
            if idx == 0:
                sec = doc.sections[0]
                if page.width > page.height:
                    sec.orientation = WD_ORIENT.LANDSCAPE
                    sec.page_width, sec.page_height = Pt(page.width), Pt(page.height)
                sec.left_margin = sec.right_margin = Pt(36)
                sec.top_margin = sec.bottom_margin = Pt(36)
            else:
                doc.add_page_break()
            tables = page.find_tables()
            table_boxes = [t.bbox for t in tables]

            def _outside(obj, boxes=table_boxes):
                return not any(b[0] <= obj["x0"] and obj["x1"] <= b[2] and b[1] <= obj["top"] and obj["bottom"] <= b[3] for b in boxes)

            text = page.filter(_outside).extract_text() if table_boxes else page.extract_text()
            for line in (text or "").splitlines():
                if line.strip():
                    doc.add_paragraph(line)
            for t in tables:
                rows = t.extract()
                if not rows:
                    continue
                ncols = max(len(r) for r in rows)
                dt = doc.add_table(rows=len(rows), cols=ncols)
                dt.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row):
                        cell = dt.cell(ri, ci)
                        cell.text = (val or "").strip()
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.size = Pt(8)
                doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
