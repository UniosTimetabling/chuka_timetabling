"""
For every Program-Year-Semester-Intake that has Specialization Stems and/or
Elective (Selection) Groups, this works out what a REAL student actually
ends up taking once they've made their picks — one stem per specialization
category, one course per elective group — and flags any such combination
where the resulting course count exceeds 9.

Why per-combination, not just "smallest possible": a student can pick ANY
stem in a category, not just the smallest one. This lists each individual
stem choice (or combination of choices, if a program-year has more than one
specialization category) as its own row, with the exact set of courses that
combination has to be collision-checked against on the timetable — i.e. the
courses that must never be scheduled in the same slot, because a student
taking that combination attends all of them.

Merged/combined course groups (see CombinedCourseGroup) still count as ONE
unit throughout, same as elsewhere in this reporting suite.

Usage:
    python manage.py check_stem_elective_load
    python manage.py check_stem_elective_load --output /path/to/file.xlsx
    python manage.py check_stem_elective_load --threshold 9
"""
import itertools
import os
from collections import defaultdict
from datetime import datetime

from django.core.management.base import BaseCommand, CommandError

from course_allocation.models import CourseAllocation, CombinedCourseGroup

try:
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    pd = None


class Command(BaseCommand):
    help = (
        "For program-years with Specialization Stems and/or Elective groups, "
        "enumerate every real stem-pick scenario and flag those where the "
        "resulting course count exceeds the threshold (default 9), listing "
        "the exact collision-check course set for each."
    )

    def add_arguments(self, parser):
        parser.add_argument("--output", type=str, default="",
                             help="Output .xlsx path. Defaults to ./exports/stem_elective_load_report_<timestamp>.xlsx")
        parser.add_argument("--threshold", type=int, default=9,
                             help="Flag a stem-pick scenario if its course count exceeds this (default 9).")

    # ------------------------------------------------------------------
    def handle(self, *args, **options):
        if pd is None:
            raise CommandError(
                "This command needs pandas + openpyxl: "
                "pip install pandas openpyxl --break-system-packages"
            )

        threshold = options["threshold"]
        out_path = options["output"] or self._default_path()
        os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)

        allocations = list(CourseAllocation.objects.select_related(
            "program", "program__department", "program__department__faculty",
            "program_course",
        ).prefetch_related(
            "specialization_stems", "specialization_stems__category", "selection_groups",
        ))

        # Merged/combined course groups collapse to ONE representative unit.
        canonical_id = {}
        for group in CombinedCourseGroup.objects.prefetch_related("allocations"):
            member_ids = list(group.allocations.values_list("id", flat=True))
            if not member_ids:
                continue
            rep = group.primary_allocation_id or member_ids[0]
            for mid in member_ids:
                canonical_id[mid] = rep

        def canon(a):
            return canonical_id.get(a.id, a.id)

        # Group by Program-Year-Semester-Intake
        py_groups = defaultdict(list)
        for a in allocations:
            if not a.program or not a.program_course:
                continue
            intake_norm = (a.intake or "").strip().capitalize()
            key = (a.program.name, a.program_course.year, a.program_course.semester, intake_norm)
            py_groups[key].append(a)

        rows = []
        for (program, year, semester, intake), allocs in py_groups.items():
            by_canon = {}
            for a in allocs:
                by_canon.setdefault(canon(a), a)
            unit_reps = list(by_canon.values())

            elective_units = [a for a in unit_reps if a.is_elective or a.selection_groups.exists()]
            stem_units = [a for a in unit_reps if a.specialization_stems.exists()]

            if not elective_units and not stem_units:
                continue  # this script only covers program-years that actually have picks to make

            dept = unit_reps[0].program.department if unit_reps[0].program.department_id else None

            shared_compulsory = [
                a for a in unit_reps
                if not a.student_group_id and not a.is_elective
                and not a.selection_groups.exists() and not a.specialization_stems.exists()
            ]
            compulsory_codes = sorted({a.course_code for a in shared_compulsory})
            compulsory_count = len(compulsory_codes)

            # Elective/selection groups: one course PICKED per group (we
            # don't know which specific one, so it's shown as a "choose one
            # of" note, but still costs exactly 1 unit per group).
            group_to_codes = defaultdict(set)
            group_names = {}
            for a in elective_units:
                for sg in a.selection_groups.all():
                    group_to_codes[sg.id].add(a.course_code)
                    group_names[sg.id] = sg.name
            ungrouped_elective_codes = sorted({
                a.course_code for a in elective_units if not a.selection_groups.exists()
            })
            elective_add = len(group_to_codes) + len(ungrouped_elective_codes)
            elective_notes = [
                f"Choose 1 of [{', '.join(sorted(codes))}] ({group_names[gid]})"
                for gid, codes in group_to_codes.items()
            ] + [f"{code} (elective)" for code in ungrouped_elective_codes]

            # Specialization stems, grouped by category -> stem -> course codes
            cat_stems = defaultdict(lambda: defaultdict(set))
            stem_names = {}
            for a in stem_units:
                for stem in a.specialization_stems.all():
                    cat_stems[stem.category_id][stem.id].add(a.course_code)
                    stem_names[stem.id] = stem.name

            if cat_stems:
                # One choice per category; enumerate every combination across
                # categories (usually just one category, but handle >1 too).
                per_category_options = [
                    list(stems.items()) for stems in cat_stems.values()  # [(stem_id, {codes}), ...]
                ]
                for combo in itertools.product(*per_category_options):
                    stem_ids = [sid for sid, _ in combo]
                    stem_code_sets = [codes for _, codes in combo]
                    stem_codes = sorted(set().union(*stem_code_sets)) if stem_code_sets else []
                    stem_count = len(stem_codes)

                    total_units = compulsory_count + stem_count + elective_add

                    if total_units > threshold:
                        picked_stem_names = ", ".join(stem_names[sid] for sid in stem_ids)
                        full_course_list = compulsory_codes + stem_codes
                        rows.append({
                            "Program": program,
                            "Department": dept.name if dept else "",
                            "Year": year, "Semester": semester, "Intake": intake,
                            "Stem(s) Picked": picked_stem_names,
                            "Total Units": total_units,
                            "Compulsory Units": compulsory_count,
                            "Stem Units": stem_count,
                            "Elective Picks": elective_add,
                            "Compulsory Courses": "; ".join(compulsory_codes),
                            "Stem Courses": "; ".join(stem_codes),
                            "Elective Picks Detail": "; ".join(elective_notes) if elective_notes else "",
                            "Full Collision-Check Course Set": "; ".join(full_course_list) + (
                                ("; " + "; ".join(elective_notes)) if elective_notes else ""
                            ),
                        })
            else:
                # No stems at all — just compulsory + elective picks.
                total_units = compulsory_count + elective_add
                if total_units > threshold:
                    rows.append({
                        "Program": program,
                        "Department": dept.name if dept else "",
                        "Year": year, "Semester": semester, "Intake": intake,
                        "Stem(s) Picked": "None",
                        "Total Units": total_units,
                        "Compulsory Units": compulsory_count,
                        "Stem Units": 0,
                        "Elective Picks": elective_add,
                        "Compulsory Courses": "; ".join(compulsory_codes),
                        "Stem Courses": "",
                        "Elective Picks Detail": "; ".join(elective_notes) if elective_notes else "",
                        "Full Collision-Check Course Set": "; ".join(compulsory_codes) + (
                            ("; " + "; ".join(elective_notes)) if elective_notes else ""
                        ),
                    })

        rows.sort(key=lambda r: -r["Total Units"])

        cols = ["Program", "Department", "Year", "Semester", "Intake", "Stem(s) Picked",
                "Total Units", "Compulsory Units", "Stem Units", "Elective Picks",
                "Compulsory Courses", "Stem Courses", "Elective Picks Detail",
                "Full Collision-Check Course Set"]
        df = pd.DataFrame(rows, columns=cols)

        sheet_name = "Stem+Elective Load"
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)

        self._style_workbook(out_path, sheet_name, cols)

        self.stdout.write(self.style.SUCCESS(
            f"Flagged {len(rows)} program-year/stem-pick combination(s) exceeding {threshold} units."
        ))
        self.stdout.write(self.style.SUCCESS(f"Wrote {out_path}"))

    # ------------------------------------------------------------------
    def _style_workbook(self, path, sheet_name, cols):
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet_name]
        header_fill = PatternFill(start_color="1E7145", end_color="1E7145", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        ws.freeze_panes = "A2"
        wide_cols = {"Compulsory Courses", "Stem Courses", "Elective Picks Detail", "Full Collision-Check Course Set"}
        for col_idx, col_name in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            width = 55 if col_name in wide_cols else (26 if col_name in ("Program", "Department", "Stem(s) Picked") else 15)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        wb.save(path)

    def _default_path(self):
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        return os.path.join("exports", f"stem_elective_load_report_{ts}.xlsx")
