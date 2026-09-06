"""
course_allocation/program_enrollment.py

Program Enrollment management (entry-year model):
  - program_enrollment_page      : main page view
  - save_program_enrollment      : create/update ONE cohort's student count
                                    (keyed by program + entry_year, applies to
                                    both semesters automatically -- there is
                                    no separate per-semester row any more)
  - save_enrollment_single_sem   : kept for backward-compatible URL name;
                                    now just edits a cohort's student count
  - delete_program_enrollment    : delete a cohort record
  - get_program_enrollments      : AJAX fetch enrollments for a program
  - advance_academic_year        : bump/set the single global reference year
                                    (replaces the old per-record "shift")
  - upload_enrollment_file       : import CSV / XLSX / JSON
  - enrollment_template          : download CSV or XLSX template
"""

import csv
import io
import json

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from core.rbac import allowed_roles, Role
from django.core.exceptions import PermissionDenied
from django.http import (
    HttpResponse,
    HttpResponseBadRequest,
    HttpResponseForbidden,
    JsonResponse,
)
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST

from department_management.models import Department
from program_management.models import Program
from .detect_user_department import detect_user_department
from .models import ProgramEnrollment, AcademicYearTracker
from .auto_allocate_courses import (
    ADMIN_GROUPS,
    ALL_ALLOWED,
    _resolve_user_dept,
    check_user_permission,
    get_homepage_url,
)

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ─────────────────────────────────────────────────────────────────────────────
# Page
# ─────────────────────────────────────────────────────────────────────────────

@allowed_roles(Role.COD, Role.COD_ADMIN, Role.SUDO)
def program_enrollment_page(request):
    if not check_user_permission(request.user, ALL_ALLOWED):
        raise PermissionDenied

    user = request.user
    is_admin, user_dept, departments = _resolve_user_dept(user)
    homepage_url = get_homepage_url(user)

    if not is_admin and not user_dept:
        messages.error(request, "Unable to detect your department. Contact admin.")
        return redirect(homepage_url)

    tracker = AcademicYearTracker.get_current()
    reference_year = tracker.current_year

    if user_dept:
        programs = Program.objects.filter(department=user_dept).order_by("name")
        enrollments = (
            ProgramEnrollment.objects
            .filter(program__department=user_dept)
            .select_related("program", "program__department")
        )
    else:
        programs = (
            Program.objects.all()
            .select_related("department")
            .order_by("department__name", "name")
        )
        enrollments = (
            ProgramEnrollment.objects.all()
            .select_related("program", "program__department")
        )

    # Attach the derived "year of study" to each row for display -- computed
    # here, not stored, so it's always in sync with the tracker.
    enrollments = list(enrollments)
    for e in enrollments:
        e.study_year = e.current_study_year(reference_year)

    return render(request, "course_allocation/program_enrollment.html", {
        "programs": programs,
        "enrollments": enrollments,
        "departments": departments,
        "user_department": user_dept,
        "is_superuser": user.is_superuser,
        "is_admin": is_admin,
        "homepage_url": homepage_url,
        "reference_year": reference_year,
    })


# ─────────────────────────────────────────────────────────────────────────────
# Save — ONE row per (program, entry_year). Applies to both semesters
# automatically since there's no semester field on this model any more.
# ─────────────────────────────────────────────────────────────────────────────

@allowed_roles(Role.COD, Role.COD_ADMIN, Role.SUDO)
@require_POST
def save_program_enrollment(request):
    if not check_user_permission(request.user, ALL_ALLOWED):
        return HttpResponseForbidden("No permission.")

    program_id   = request.POST.get("program_id")
    entry_year   = request.POST.get("entry_year")
    num_students = request.POST.get("number_of_students", 0)

    if not all([program_id, entry_year]):
        return JsonResponse(
            {"success": False, "error": "Program and entry year are required."},
            status=400,
        )

    try:
        program      = Program.objects.select_related("department").get(id=program_id)
        entry_year   = int(entry_year)
        num_students = int(num_students)
    except (Program.DoesNotExist, ValueError):
        return JsonResponse({"success": False, "error": "Invalid data."}, status=400)

    if not check_user_permission(request.user, ADMIN_GROUPS):
        user_dept = detect_user_department(request.user)
        if user_dept and program.department_id != user_dept.id:
            return HttpResponseForbidden("You can only manage enrollments for your own department.")

    enrollment, created = ProgramEnrollment.objects.update_or_create(
        program=program, entry_year=entry_year,
        defaults={"number_of_students": num_students},
    )

    return JsonResponse({
        "success": True,
        "enrollment_id": enrollment.id,
        "program": program.name,
        "program_id": program.id,
        "department": program.department.name,
        "department_id": program.department_id,
        "entry_year": entry_year,
        "study_year": enrollment.current_study_year(),
        "number_of_students": num_students,
        "created": created,
    })


# ─────────────────────────────────────────────────────────────────────────────
# Save single record (inline table edit) — kept under its old URL name for
# backward compatibility with the frontend; there's only ever one row per
# cohort now, so this just edits that row's student count.
# ─────────────────────────────────────────────────────────────────────────────

@allowed_roles(Role.COD, Role.COD_ADMIN, Role.SUDO)
@require_POST
def save_enrollment_single_sem(request):
    """Update one cohort's student count in-place."""
    if not check_user_permission(request.user, ALL_ALLOWED):
        return HttpResponseForbidden("No permission.")

    enrollment_id = request.POST.get("enrollment_id")
    num_students  = request.POST.get("number_of_students", 0)

    try:
        e = ProgramEnrollment.objects.select_related("program__department").get(id=enrollment_id)
        num_students = int(num_students)
    except (ProgramEnrollment.DoesNotExist, ValueError):
        return JsonResponse({"success": False, "error": "Record not found."}, status=404)

    if not check_user_permission(request.user, ADMIN_GROUPS):
        user_dept = detect_user_department(request.user)
        if user_dept and e.program.department_id != user_dept.id:
            return HttpResponseForbidden("No permission.")

    e.number_of_students = num_students
    e.save(update_fields=["number_of_students"])

    return JsonResponse({"success": True, "enrollment_id": e.id, "number_of_students": num_students})


# ─────────────────────────────────────────────────────────────────────────────
# Delete
# ─────────────────────────────────────────────────────────────────────────────

@allowed_roles(Role.COD, Role.COD_ADMIN, Role.SUDO)
@require_POST
def delete_program_enrollment(request):
    if not check_user_permission(request.user, ALL_ALLOWED):
        return HttpResponseForbidden("No permission.")

    enrollment_id = request.POST.get("enrollment_id")
    try:
        e = ProgramEnrollment.objects.select_related("program__department").get(id=enrollment_id)
    except ProgramEnrollment.DoesNotExist:
        return JsonResponse({"success": False, "error": "Record not found."}, status=404)

    if not check_user_permission(request.user, ADMIN_GROUPS):
        user_dept = detect_user_department(request.user)
        if not user_dept or e.program.department_id != user_dept.id:
            return HttpResponseForbidden("No permission.")

    e.delete()
    return JsonResponse({"success": True, "deleted": 1})


# ─────────────────────────────────────────────────────────────────────────────
# Get enrollments for a program (AJAX)
# ─────────────────────────────────────────────────────────────────────────────

@allowed_roles(Role.COD, Role.COD_ADMIN, Role.SUDO)
@require_GET
def get_program_enrollments(request):
    program_id = request.GET.get("program_id")
    if not program_id:
        return JsonResponse({"success": False, "error": "program_id required."}, status=400)

    reference_year = AcademicYearTracker.get_current().current_year
    enrollments = ProgramEnrollment.objects.filter(program_id=program_id).order_by("-entry_year")

    data = [
        {
            "id": e.id,
            "entry_year": e.entry_year,
            "study_year": e.current_study_year(reference_year),
            "number_of_students": e.number_of_students,
        }
        for e in enrollments
    ]
    return JsonResponse({"success": True, "reference_year": reference_year, "enrollments": data})


# ─────────────────────────────────────────────────────────────────────────────
# ADVANCE ACADEMIC YEAR
#   Replaces the old per-record "shift forward/reverse". Nothing on
#   ProgramEnrollment changes at all -- bumping this ONE global number
#   instantly reclassifies every cohort's year-of-study everywhere (COD
#   panel, program-enrollment page, auto-allocation) because study_year is
#   always computed as current_year - entry_year + 1, never stored.
# ─────────────────────────────────────────────────────────────────────────────

@allowed_roles(Role.COD, Role.COD_ADMIN, Role.SUDO)
@require_POST
def advance_academic_year(request):
    if not check_user_permission(request.user, ALL_ALLOWED):
        return HttpResponseForbidden("No permission.")

    # Only admins may set this globally -- it affects every department at once.
    if not check_user_permission(request.user, ADMIN_GROUPS):
        return HttpResponseForbidden("Only an admin can advance the institution's academic year.")

    direction = request.POST.get("direction")   # "forward" | "reverse" | None
    explicit_year = request.POST.get("year")    # optional explicit override

    tracker = AcademicYearTracker.get_current()

    if explicit_year:
        try:
            new_year = int(explicit_year)
        except ValueError:
            return HttpResponseBadRequest("year must be an integer.")
    elif direction == "reverse":
        new_year = tracker.current_year - 1
    else:
        new_year = tracker.current_year + 1

    tracker = AcademicYearTracker.set_current_year(new_year)

    return JsonResponse({
        "success": True,
        "previous_year": tracker.current_year if explicit_year else (
            new_year - 1 if direction != "reverse" else new_year + 1
        ),
        "current_year": tracker.current_year,
    })


# Backward-compatible alias -- old templates/JS may still call the old name.
shift_enrollment_year = advance_academic_year


# ─────────────────────────────────────────────────────────────────────────────
# FILE UPLOAD  — CSV / JSON / XLSX
# Required columns: program_name (or program_id), entry_year, number_of_students
# ─────────────────────────────────────────────────────────────────────────────

@allowed_roles(Role.COD, Role.COD_ADMIN, Role.SUDO)
@require_POST
def upload_enrollment_file(request):
    if not check_user_permission(request.user, ALL_ALLOWED):
        return HttpResponseForbidden("No permission.")

    uploaded = request.FILES.get("enrollment_file")
    if not uploaded:
        return JsonResponse({"success": False, "error": "No file uploaded."}, status=400)

    if uploaded.size > 5 * 1024 * 1024:
        return JsonResponse({"success": False, "error": "File exceeds 5 MB limit."}, status=400)

    ext  = uploaded.name.rsplit(".", 1)[-1].lower()
    rows = []

    try:
        if ext == "csv":
            decoded = uploaded.read().decode("utf-8-sig")
            rows    = list(csv.DictReader(io.StringIO(decoded)))

        elif ext == "json":
            decoded = uploaded.read().decode("utf-8")
            parsed  = json.loads(decoded)
            rows    = parsed if isinstance(parsed, list) else parsed.get("enrollments", [])

        elif ext in ("xlsx", "xls"):
            if not HAS_OPENPYXL:
                return JsonResponse(
                    {"success": False, "error": "openpyxl not installed on server. Please upload a CSV file."},
                    status=400,
                )
            wb      = openpyxl.load_workbook(uploaded, read_only=True, data_only=True)
            ws      = wb.active
            headers = [
                str(c.value).strip().lower() if c.value else ""
                for c in next(ws.iter_rows(min_row=1, max_row=1))
            ]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if any(v is not None for v in row):
                    rows.append(dict(zip(headers, row)))
        else:
            return JsonResponse(
                {"success": False, "error": f"Unsupported file type: .{ext}. Use CSV, XLSX, or JSON."},
                status=400,
            )
    except Exception as e:
        return JsonResponse({"success": False, "error": f"Could not parse file: {e}"}, status=400)

    user_dept = None
    if not check_user_permission(request.user, ADMIN_GROUPS):
        user_dept = detect_user_department(request.user)

    saved   = 0
    skipped = 0
    errors  = []

    for i, row in enumerate(rows, start=1):
        row = {str(k).strip().lower(): v for k, v in row.items() if k}
        try:
            # Resolve program
            prog = None
            pid_val = row.get("program_id") or row.get("id")
            if pid_val:
                try:
                    prog = Program.objects.select_related("department").get(id=int(str(pid_val).strip()))
                except (Program.DoesNotExist, ValueError):
                    pass
            if not prog:
                name_val = str(row.get("program_name") or row.get("program") or "").strip()
                if name_val:
                    prog = Program.objects.filter(
                        name__iexact=name_val
                    ).select_related("department").first()

            if not prog:
                errors.append(f"Row {i}: program not found — '{row.get('program_name', '?')}'")
                skipped += 1
                continue

            if user_dept and prog.department_id != user_dept.id:
                errors.append(f"Row {i}: {prog.name} is not in your department.")
                skipped += 1
                continue

            entry_year_val = (
                row.get("entry_year") or row.get("intake_year")
                or row.get("year") or timezone.now().year
            )
            entry_year = int(str(entry_year_val).strip())
            stu_val    = row.get("number_of_students") or row.get("students") or 0
            students   = int(str(stu_val).strip())

            ProgramEnrollment.objects.update_or_create(
                program=prog, entry_year=entry_year,
                defaults={"number_of_students": students},
            )
            saved += 1
        except Exception as e:
            errors.append(f"Row {i}: {e}")
            skipped += 1

    return JsonResponse({
        "success": True,
        "saved":   saved,
        "skipped": skipped,
        "errors":  errors[:25],
    })


# ─────────────────────────────────────────────────────────────────────────────
# TEMPLATE DOWNLOAD  — CSV or XLSX
# ─────────────────────────────────────────────────────────────────────────────

@allowed_roles(Role.COD, Role.COD_ADMIN, Role.SUDO)
@require_GET
def enrollment_template(request):
    """
    Downloadable CSV/XLSX that the COD can edit and re-upload -- upload is
    upsert (update_or_create on program+entry_year), so re-uploading this
    file always "adds or updates", never duplicates.

    Scope: a COD gets every program in their own department; an admin/
    superuser gets every program in the system.

    Content: EVERY existing ProgramEnrollment row for those programs (real
    numbers, not zeros), arranged by entry_year (newest first), then program
    name within each year -- this is the "update" case.

    Optional ?new_entry_year=YYYY: additionally guarantees a row for that
    entry year for EVERY program in scope (0 students where one doesn't
    already exist) so the COD can fill in a brand-new intake for every
    program in one file -- this is the "add a new enrollment year" case.
    The frontend asks the user first and, if they say yes, appends this
    param with the year they choose.
    """
    fmt            = request.GET.get("fmt", "csv")
    new_entry_year = request.GET.get("new_entry_year", "").strip()

    is_admin = check_user_permission(request.user, ADMIN_GROUPS)
    if is_admin:
        programs = list(Program.objects.select_related("department").order_by("department__name", "name"))
    else:
        user_dept = detect_user_department(request.user)
        programs = list(Program.objects.filter(department=user_dept).order_by("name")) if user_dept else []

    if not programs:
        return HttpResponseBadRequest("No programs found for your account/department.")

    program_ids = [p.id for p in programs]
    existing = {
        (e.program_id, e.entry_year): e.number_of_students
        for e in ProgramEnrollment.objects.filter(program_id__in=program_ids)
    }

    entry_years = {ey for (_, ey) in existing.keys()}

    new_year_int = None
    if new_entry_year:
        try:
            new_year_int = int(new_entry_year)
            entry_years.add(new_year_int)
        except ValueError:
            new_year_int = None

    reference_year = AcademicYearTracker.get_current().current_year
    if not entry_years:
        # Nothing on file yet and no new year requested -- fall back to the
        # current reference year so the file isn't empty.
        entry_years.add(reference_year)

    entry_years = sorted(entry_years, reverse=True)   # newest year first

    # Every program in scope must appear in the file at least once -- even
    # one with zero enrollment history so far -- so nothing is silently
    # missing from what the COD sees and can fill in.
    programs_with_data = {pid for (pid, _) in existing.keys()}
    fallback_year_for_missing = new_year_int or reference_year

    headers = ["program_name", "entry_year", "number_of_students"]
    rows = []
    for ey in entry_years:
        for p in programs:
            key = (p.id, ey)
            if key in existing:
                rows.append([p.name, ey, existing[key]])
            elif ey == new_year_int:
                # New-year mode guarantees a row for every program, even
                # ones with no enrollment history at all.
                rows.append([p.name, ey, 0])
            # else: this program simply never had an entry for that older
            # year -- don't fabricate a row for it.

    for p in programs:
        if p.id not in programs_with_data and new_year_int is None:
            # Program has no enrollment rows at all, and we're not already
            # in new-year mode (which already gave it a row above) -- add
            # one placeholder row so its name still shows up in the file.
            rows.append([p.name, fallback_year_for_missing, 0])

    if fmt == "csv":
        # NOTE: a plain CSV has no concept of "locked" cells -- it's just
        # text, so there is no way to stop someone editing the program_name
        # or entry_year columns in this format. Use XLSX (below) if you want
        # those columns actually protected.
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="enrollment_template.csv"'
        writer = csv.writer(response)
        writer.writerow(headers)
        for r in rows:
            writer.writerow(r)
        return response

    if fmt == "xlsx":
        if not HAS_OPENPYXL:
            return HttpResponse("openpyxl not installed.", status=400)
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Enrollment"

        EDITABLE_FILL = PatternFill("solid", fgColor="FFF9C4")   # light yellow = "edit me"
        LOCKED_FILL   = PatternFill("solid", fgColor="F3F3F3")   # light grey  = "don't edit"

        ws.append(headers)
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font       = Font(bold=True, color="FFFFFF", size=12)
            cell.fill       = PatternFill("solid", fgColor="2E7D32")
            cell.alignment  = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        for r in rows:
            ws.append(r)

        # Colour program_name (col A) and entry_year (col B) grey and
        # number_of_students (col C) yellow as a VISUAL guide only.
        # NOTE: this used to also apply real cell/sheet protection
        # (locked=True/False + ws.protection.sheet=True), but that made
        # every cell -- including the number_of_students column that's
        # supposed to be editable -- show up as protected/read-only in
        # several spreadsheet apps (Google Sheets in particular doesn't
        # honour per-cell "locked=False" inside a protected sheet the
        # way desktop Excel does). Since this was only ever meant as an
        # accidental-edit guard rather than a real security boundary,
        # we now rely on colour alone and leave every cell editable.
        last_row = ws.max_row
        for row_idx in range(2, last_row + 1):
            ws.cell(row=row_idx, column=1).fill = LOCKED_FILL
            ws.cell(row=row_idx, column=2).fill = LOCKED_FILL
            ws.cell(row=row_idx, column=3).fill = EDITABLE_FILL

        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 24

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = HttpResponse(
            buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="enrollment_template.xlsx"'
        return response

    return HttpResponseBadRequest("Unsupported format. Use ?fmt=csv or ?fmt=xlsx")